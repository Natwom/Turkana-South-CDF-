from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

EXPORT_HEADERS = ["Application Number", "Student Name", "Reg/Admission No", "ID Number",
    "NEMIS No", "Telephone", "Gender", "Date of Birth", "Constituency", "Ward",
    "Location", "Sub-Location", "Village", "Institution", "School Paybill", "School Account No",
    "Campus", "Level", "Course", "Mode of Study", "Class/Year", "Expected Completion",
    "Family Status", "Father Name", "Father Occupation", "Mother Name", "Mother Occupation",
    "Siblings Count", "Amount Requested", "Status", "Application Date"]

ALLOC_HEADERS = ["Application Number", "Student Name", "Reg/Admission No", "Institution",
    "School Paybill", "School Account No", "Course", "Education Level", "Ward", "Family Status",
    "Amount Requested", "Amount Allocated", "Allocation Status", "Allocation Date",
    "Approved By", "Remarks"]

def _autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 45)

def export_applications(apps) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Applications"
    ws.append(EXPORT_HEADERS)
    for c in ws[1]: c.font = Font(bold=True)
    for app in apps:
        a, fam = app.applicant, app.family
        ws.append([app.application_number, a.full_name if a else "", a.reg_number if a else "",
            a.id_number if a else "", a.nemis_number if a else "", a.telephone if a else "",
            a.gender if a else "", str(a.dob) if a and a.dob else "",
            a.constituency if a else "", a.ward if a else "", a.location if a else "",
            a.sub_location if a else "", a.village if a else "", a.institution if a else "",
            a.school_paybill if a else "", a.school_account_number if a else "",
            a.campus if a else "", a.level_of_study if a else "", a.course if a else "",
            a.mode_of_study if a else "", a.class_year if a else "",
            a.expected_completion if a else "", app.family_status or "",
            fam.father.name if fam and fam.father else "", fam.father.occupation if fam and fam.father else "",
            fam.mother.name if fam and fam.mother else "", fam.mother.occupation if fam and fam.mother else "",
            len(app.siblings), float(app.amount_requested or 0), app.status,
            str(app.created_at)])
    _autosize(ws)
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

def export_allocations(apps, period) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Allocations"
    ws.append(ALLOC_HEADERS)
    for c in ws[1]: c.font = Font(bold=True)
    total_req = total_alloc = 0.0
    for app in apps:
        a = app.applicant
        alloc = app.allocation
        total_req += float(app.amount_requested or 0)
        amt = float(alloc.amount) if alloc else 0
        total_alloc += amt
        ws.append([app.application_number, a.full_name if a else "", a.reg_number if a else "",
            a.institution if a else "", a.school_paybill if a else "", a.school_account_number if a else "",
            a.course if a else "", a.level_of_study if a else "",
            a.ward if a else "", app.family_status or "", float(app.amount_requested or 0), amt,
            "Allocated" if alloc else "Pending", str(alloc.created_at) if alloc else "",
            alloc.allocated_by if alloc else "", alloc.remarks if alloc else ""])
    ws.append([]); ws.append(["", "", "", "", "", "", "", "", "", "TOTAL REQUESTED", total_req])
    ws.append(["", "", "", "", "", "", "", "", "", "TOTAL ALLOCATED", total_alloc])
    ws.append(["", "", "", "", "", "", "", "", "", "REMAINING FUND",
               float(period.total_fund or 0) - total_alloc])
    _autosize(ws)
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

def allocation_template() -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Allocation Template"
    ws.append(["Application Number", "Student Name", "Reg/Admission No", "Institution",
               "Ward", "Course", "Amount Requested", "Amount Allocated", "Status", "Remarks"])
    for c in ws[1]: c.font = Font(bold=True)
    _autosize(ws)
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

def parse_allocation_import(content: bytes):
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    errors, rows = [], []
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    expected = ["Application Number", "Amount Allocated"]
    for e in expected:
        if e not in headers:
            errors.append(f"Missing required column: {e}")
    if errors: return errors, rows
    i_app, i_amt = headers.index("Application Number"), headers.index("Amount Allocated")
    seen = set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[i_app]: continue
        app_no = str(row[i_app]).strip()
        if app_no in seen: errors.append(f"Row {idx}: duplicate application number {app_no}")
        seen.add(app_no)
        try:
            amt = float(row[i_amt])
            if amt < 0: raise ValueError
        except (TypeError, ValueError):
            errors.append(f"Row {idx}: invalid amount '{row[i_amt]}'"); continue
        rows.append({"application_number": app_no, "amount": amt,
                     "remarks": str(row[headers.index("Remarks")] or "") if "Remarks" in headers else ""})
    return errors, rows